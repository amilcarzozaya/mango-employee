"""MANGO Meeting Intelligence document renderers.

Only validated structured reports reach these functions. Word/PDF are optional
dependencies; ordinary MANGO installation continues to work without them.
"""
from __future__ import annotations
import html
import json
from pathlib import Path

def _safe(value):
    return str(value if value is not None and value != "" else "NO DEFINIDO").replace("|", "\\|").replace("\n", " ")

def sections(report):
    def ev(x):
        return "Evidencia: «" + x["source_excerpt"] + "»"
    return [
        ("Resumen ejecutivo", [report["executive_summary"]]),
        ("Tres puntos críticos (máximo)", [
            x["title"] + " [" + str(x["priority_score"]) + "/12] — " +
            x["reason"] + ". " + ev(x) for x in report["critical_points"]]),
        ("Compromisos", [
            x["description"] + " — Responsable: " + (x["owner"] or "NO DEFINIDO") +
            "; fecha: " + (x["due_date"] or x["due_text"] or "NO DEFINIDA") +
            "; estado: " + x["status"] + ". " + ev(x)
            for x in report["commitments"]]),
        ("Decisiones", [x["description"] + ". " + ev(x) for x in report["decisions"]]),
        ("Pendientes", [
            x["description"] + " — Responsable: " + (x["owner"] or "NO DEFINIDO") +
            ". " + ev(x) for x in report["pending"]]),
        ("Por confirmar", report["review_required"]),
    ]

def markdown_report(report):
    m=report["meeting"]
    lines=[
        "# MANGO Meeting Intelligence",
        "",
        "**Reunión:** "+m["title"],
        "**Fecha:** "+(m["date"] or "NO DEFINIDA"),
        "**Zona horaria:** "+m["timezone"],
        "**Participantes:** "+(", ".join(m["participants"]) or "NO DEFINIDOS"),
        "**Objetivo:** "+(m["objective"] or "NO DEFINIDO"),
        "",
    ]
    lines += ["## Resumen ejecutivo", "", report["executive_summary"], "",
              "## Tres puntos críticos (máximo)", ""]
    for item in report["critical_points"]:
        lines += ["### "+item["title"]+" — "+str(item["priority_score"])+"/12",
                  item["reason"],
                  "**Criterios:** "+", ".join(k+"="+str(v) for k,v in item["criteria"].items()),
                  "**Evidencia:** «"+item["source_excerpt"]+"»", ""]
    if not report["critical_points"]:
        lines += ["No se identificaron asuntos críticos con evidencia suficiente.", ""]
    lines += [
        "## Tareas y fechas", "",
        "| ID | Tarea | Responsable | Fecha | Estado | Evidencia |",
        "|---|---|---|---|---|---|",
    ]
    for x in report["tasks"]:
        lines.append(
            "| "+x["id"]+" | "+_safe(x["description"])+" | "+_safe(x["owner"])+
            " | "+_safe(x["due_date"] or x["due_text"])+" | "+x["status"]+
            " | "+_safe(x["source_excerpt"])+" |")
    if not report["tasks"]:
        lines.append("| — | Sin tareas verificables | — | — | — | — |")
    lines += [""]
    for heading, items in sections(report)[2:]:
        lines += ["## "+heading, ""]
        lines += ["- "+x for x in items] if items else ["Sin elementos verificables."]
        lines.append("")
    lines += [
        "---", "Borrador para revisión humana. No se enviaron mensajes, "
        "no se asignaron tareas en servicios externos y no se promovió memoria automáticamente.",
        "",
    ]
    return "\n".join(lines)

def docx_report(report, path):
    try:
        from docx import Document
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise ValueError("Falta python-docx: pip install -e '.[meeting]'") from exc
    doc=Document()
    page=doc.sections[0]
    page.top_margin=Inches(.7)
    page.bottom_margin=Inches(.7)
    doc.styles["Normal"].font.size=Pt(10)
    doc.add_heading("MANGO Meeting Intelligence",0)
    m=report["meeting"]
    doc.add_paragraph(m["title"],style="Subtitle")
    doc.add_paragraph(
        "Fecha: "+(m["date"] or "NO DEFINIDA")+" | Zona: "+m["timezone"]+
        " | Participantes: "+(", ".join(m["participants"]) or "NO DEFINIDOS"))
    doc.add_paragraph("Objetivo: "+(m["objective"] or "NO DEFINIDO"))
    doc.add_heading("Resumen ejecutivo",1)
    doc.add_paragraph(report["executive_summary"])
    doc.add_heading("Tres puntos críticos (máximo)",1)
    if not report["critical_points"]:
        doc.add_paragraph("Sin puntos críticos fundamentados.")
    for x in report["critical_points"]:
        doc.add_heading(x["title"]+" ("+str(x["priority_score"])+"/12)",2)
        doc.add_paragraph(x["reason"])
        doc.add_paragraph("Evidencia: «"+x["source_excerpt"]+"»")
    doc.add_heading("Tareas y fechas",1)
    table=doc.add_table(rows=1, cols=5)
    table.style="Light Shading Accent 1"
    for i,name in enumerate(("Tarea","Responsable","Fecha","Estado","Evidencia")):
        table.rows[0].cells[i].text=name
    for x in report["tasks"]:
        cells=table.add_row().cells
        for i,value in enumerate((
            x["description"], x["owner"] or "NO DEFINIDO",
            x["due_date"] or x["due_text"] or "NO DEFINIDA",
            x["status"], x["source_excerpt"])):
            cells[i].text=value
    if not report["tasks"]:
        cells=table.add_row().cells
        cells[0].text="Sin tareas verificables"
    for heading,items in sections(report)[2:]:
        doc.add_heading(heading,1)
        if not items:
            doc.add_paragraph("Sin elementos verificables.")
        for item in items:
            doc.add_paragraph(item,style="List Bullet")
    doc.add_paragraph(
        "Borrador para revisión humana; no autoriza envíos externos.")
    doc.save(str(path))

def pdf_report(report, path):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except ImportError as exc:
        raise ValueError("Falta reportlab: pip install -e '.[meeting]'") from exc
    regular=Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    if regular.is_file() and "MangoUnicode" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("MangoUnicode",str(regular)))
    name="MangoUnicode" if "MangoUnicode" in pdfmetrics.getRegisteredFontNames() else "Helvetica"
    styles=getSampleStyleSheet()
    for k in ("Title","Heading1","Heading2","Normal","BodyText"):
        styles[k].fontName=name
    styles["BodyText"].leading=15
    def para(value,style="BodyText"):
        return Paragraph(html.escape(str(value)),styles[style])
    story=[para("MANGO Meeting Intelligence","Title")]
    m=report["meeting"]
    story.extend([
        para(m["title"],"Heading1"),
        para("Fecha: "+(m["date"] or "NO DEFINIDA")+" | Zona: "+m["timezone"]),
        para("Participantes: "+(", ".join(m["participants"]) or "NO DEFINIDOS")),
        para("Objetivo: "+(m["objective"] or "NO DEFINIDO")),
    ])
    story.extend([Spacer(1,12),para("Resumen ejecutivo","Heading1"),
                  para(report["executive_summary"])])
    story.extend([Spacer(1,12),para("Tres puntos críticos (máximo)","Heading1")])
    if not report["critical_points"]:
        story.append(para("Sin puntos críticos fundamentados."))
    for x in report["critical_points"]:
        story += [para(x["title"]+" ("+str(x["priority_score"])+"/12)","Heading2"),
                  para(x["reason"]),para("Evidencia: «"+x["source_excerpt"]+"»"),
                  Spacer(1,5)]
    story += [Spacer(1,10),para("Tareas y fechas","Heading1")]
    if not report["tasks"]:
        story.append(para("Sin tareas verificables."))
    for x in report["tasks"]:
        story += [
            para(x["description"],"Heading2"),
            para("Responsable: "+(x["owner"] or "NO DEFINIDO")+
                 " | Fecha: "+(x["due_date"] or x["due_text"] or "NO DEFINIDA")+
                 " | Estado: "+x["status"]),
            para("Evidencia: «"+x["source_excerpt"]+"»"),
            Spacer(1,5),
        ]
    for heading,items in sections(report)[2:]:
        story += [Spacer(1,10),para(heading,"Heading1")]
        story += [para("• "+item) for item in items] if items else [
            para("Sin elementos verificables.")]
    story += [Spacer(1,12),para(
        "Borrador para revisión humana; no autoriza envíos externos.")]
    SimpleDocTemplate(
        str(path),pagesize=A4,leftMargin=48,rightMargin=48,
        topMargin=44,bottomMargin=44).build(story)

def export_report(report, directory, formats=("json","md")):
    supported={"json","md","docx","pdf"}
    formats=tuple(dict.fromkeys(formats))
    if not formats or any(f not in supported for f in formats):
        raise ValueError("Formatos válidos: json,md,docx,pdf.")
    # Preflight dependencies before creating any file to avoid partial export.
    if "docx" in formats:
        try:
            import docx  # noqa: F401
        except ImportError as exc:
            raise ValueError("Instala Word/PDF: pip install -e '.[meeting]'") from exc
    if "pdf" in formats:
        try:
            import reportlab  # noqa: F401
        except ImportError as exc:
            raise ValueError("Instala Word/PDF: pip install -e '.[meeting]'") from exc
    out=Path(directory)
    out.mkdir(parents=True,exist_ok=True)
    paths={fmt:out/(report["report_id"]+"."+fmt) for fmt in formats}
    if any(p.exists() for p in paths.values()):
        raise ValueError("Un archivo de reporte ya existe; no se sobrescribe.")
    for fmt,path in paths.items():
        if fmt=="json":
            path.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",
                            encoding="utf-8")
        elif fmt=="md":
            path.write_text(markdown_report(report),encoding="utf-8")
        elif fmt=="docx":
            docx_report(report,path)
        else:
            pdf_report(report,path)
    return {fmt:str(path) for fmt,path in paths.items()}
