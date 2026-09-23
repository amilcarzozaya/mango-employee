"""Shared quotation document renderers: no independent monetary arithmetic."""
from __future__ import annotations

from pathlib import Path
import html
import json
import os
import hashlib

FORMATS = {"json", "md", "docx", "pdf"}


def preflight_formats(formats):
    formats = tuple(formats)
    if not formats or any(f not in FORMATS for f in formats) or len(formats) != len(set(formats)):
        raise ValueError("Formatos válidos, sin duplicados: json,md,docx,pdf.")
    if "docx" in formats:
        try:
            import docx  # noqa
        except ImportError as exc:
            raise ValueError("Falta soporte Word: pip install -e '.[quote]'") from exc
    if "pdf" in formats:
        try:
            import reportlab  # noqa
        except ImportError as exc:
            raise ValueError("Falta soporte PDF: pip install -e '.[quote]'") from exc


def _money(value, quote):
    places = quote["currency_decimals"]
    from decimal import Decimal
    return quote["currency"] + " " + f"{Decimal(value):,.{places}f}"


def _header(quote):
    issuer = quote["issuer_snapshot"]["seller"]
    return [
        f"# {'COTIZACIÓN' if quote['status'] == 'issued' else 'BORRADOR DE COTIZACIÓN'}",
        "",
        f"**Folio:** {quote['folio'] or 'POR ASIGNAR (BORRADOR)'}",
        f"**Fecha:** {quote['quote_date']} · **Vigencia:** {quote['valid_until']}",
        f"**Revisión de:** {quote['revision_of'] or 'Ninguna'}",
        "",
        f"**Emisor:** {issuer['company_name']}",
        f"**Razón social:** {issuer['legal_name']}",
        f"**Contacto:** {issuer['contact_name']}",
        f"**Email:** {issuer['email']} · **Teléfono:** {issuer['phone']}",
        f"**Dirección:** {issuer['address']}",
        f"**Identificador fiscal:** {issuer['tax_id'] or 'No proporcionado'}",
        "",
        f"**Cliente:** {quote['client']['name']}",
        f"**Empresa:** {quote['client']['company_name'] or 'No proporcionada'}",
        f"**Email:** {quote['client']['email'] or 'No proporcionado'}",
        f"**Dirección:** {quote['client']['address'] or 'No proporcionada'}",
        "",
    ]


def markdown_quote(quote):
    lines = _header(quote)
    lines += [
        "## Conceptos",
        "",
        "| # | Descripción | Cantidad | Unidad | Precio unitario | Importe | Descuento | Base | Impuestos | Retenciones | Total |",
        "|---:|---|---:|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    clean = lambda s: str(s).replace("|", r"\|").replace("\n", " ")
    for line in quote["items"]:
        lines.append(
            f"| {line['line']} | {clean(line['description'])} | {line['quantity']} | "
            f"{clean(line['unit'])} | {_money(line['unit_price'],quote)} | "
            f"{_money(line['gross'],quote)} | {_money(line['discount'],quote)} | "
            f"{_money(line['taxable_base'],quote)} | {_money(line['tax_added'],quote)} | "
            f"{_money(line['tax_withheld'],quote)} | {_money(line['line_total'],quote)} |"
        )
    totals = quote["totals"]
    lines += [
        "", "## Resumen monetario",
        f"- Importe bruto: {_money(totals['gross'],quote)}",
        f"- Descuentos: −{_money(totals['discount'],quote)}",
        f"- Base después de descuentos: {_money(totals['taxable_base'],quote)}",
    ]
    for tax in quote["tax_breakdown"]:
        label = "Impuesto adicional" if tax["kind"] == "add" else "Retención"
        lines.append(
            f"- {label} — {tax['label']} ({tax['code']}, {tax['rate']}): "
            + ("−" if tax["kind"] == "withhold" else "+")
            + _money(tax["amount"], quote)
        )
    lines += [
        f"- **TOTAL:** **{_money(totals['total'], quote)}**",
        "", "## Condiciones",
        f"- Precios incluyen impuestos configurados: {'Sí' if quote['prices_include_tax'] else 'No'}",
        f"- Vigencia: {quote['valid_until']}",
        f"- Condiciones de pago: {quote['terms']['payment_terms'] or 'Por definir'}",
        f"- Entrega: {quote['terms']['delivery_terms'] or 'Por definir'}",
        f"- Notas: {quote['terms']['notes'] or 'Ninguna'}",
    ]
    if quote["exclusions"]:
        lines += ["", "## Exclusiones"] + [f"- {x}" for x in quote["exclusions"]]
    lines += [
        "", "---",
        quote["notice"],
        "No implica envío al cliente. La tasa y el tratamiento fiscal fueron configurados por el emisor; confirme su aplicabilidad.",
    ]
    if quote["status"] == "issued":
        lines += [f"Emitida: {quote['issued_at']}", f"Aprobación declarada por: {quote['approved_by']}"]
    return "\n".join(lines) + "\n"


def docx_quote(quote, path):
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise ValueError("Instala Word: pip install -e '.[quote]'") from exc
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.size = Pt(9)
    seller = quote["issuer_snapshot"]["seller"]
    heading = "COTIZACIÓN" if quote["status"] == "issued" else "BORRADOR DE COTIZACIÓN"
    doc.add_heading(heading, 0)
    doc.add_paragraph(quote["folio"] or "Sin folio — borrador", style="Subtitle")
    doc.add_paragraph(f"Fecha: {quote['quote_date']} | Válida hasta: {quote['valid_until']}")
    doc.add_heading("Emisor", level=2)
    doc.add_paragraph(
        f"{seller['company_name']} | {seller['legal_name']}\n"
        f"{seller['contact_name']} · {seller['email']} · {seller['phone']}\n{seller['address']}"
    )
    doc.add_heading("Cliente", level=2)
    cli = quote["client"]
    doc.add_paragraph(
        cli["name"] + (f" | {cli['company_name']}" if cli["company_name"] else "") +
        (f"\n{cli['email']}" if cli["email"] else "")
    )
    doc.add_heading("Conceptos", level=2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    for col, label in zip(table.rows[0].cells, (
        "Descripción", "Cant.", "P. unitario", "Descuento", "Impuestos netos", "Total"
    )):
        col.text = label
    for item in quote["items"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, (
            item["description"] + " (" + item["unit"] + ")",
            item["quantity"],
            _money(item["unit_price"], quote),
            _money(item["discount"], quote),
            _money(item["tax_added"], quote) + " − " + _money(item["tax_withheld"], quote),
            _money(item["line_total"], quote),
        )):
            cell.text = value
    doc.add_heading("Resumen", level=2)
    t = quote["totals"]
    for label, value in (
        ("Subtotal bruto", t["gross"]), ("Descuentos", t["discount"]),
        ("Base", t["taxable_base"]), ("Impuestos adicionales", t["tax_added"]),
        ("Retenciones", t["tax_withheld"]), ("TOTAL", t["total"])
    ):
        p = doc.add_paragraph()
        if label == "TOTAL":
            p.add_run(label + ": " + _money(value, quote)).bold = True
        else:
            p.add_run(label + ": " + _money(value, quote))
    for rule in quote["tax_breakdown"]:
        doc.add_paragraph(
            ("+" if rule["kind"] == "add" else "−") + rule["label"] +
            " [" + rule["code"] + "] " + _money(rule["amount"], quote)
        )
    doc.add_heading("Condiciones", level=2)
    for label, value in (
        ("Pago", quote["terms"]["payment_terms"]),
        ("Entrega", quote["terms"]["delivery_terms"]),
        ("Notas", quote["terms"]["notes"]),
    ):
        doc.add_paragraph(label + ": " + (value or "Por definir"))
    if quote["exclusions"]:
        doc.add_heading("Exclusiones", level=2)
        for exclusion in quote["exclusions"]:
            doc.add_paragraph(exclusion, style="List Bullet")
    doc.add_paragraph(quote["notice"])
    doc.add_paragraph("No implica envío al cliente ni sustituye revisión fiscal.")
    if quote["status"] == "issued":
        doc.add_paragraph("Aprobación declarada por: " + quote["approved_by"])
    footer = section.footer.paragraphs[0]
    footer.text = "MANGO Quote Builder | Cotización comercial, no comprobante fiscal"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(path))


def pdf_quote(quote, path):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
    except ImportError as exc:
        raise ValueError("Instala PDF: pip install -e '.[quote]'") from exc
    fonts = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/Library/Fonts/Arial Unicode.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    font = "Helvetica"
    for p in fonts:
        if p.is_file():
            if "MangoQuoteUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("MangoQuoteUnicode", str(p)))
            font = "MangoQuoteUnicode"
            break
    styles = getSampleStyleSheet()
    for style in ("Title", "Heading1", "Heading2", "BodyText", "Normal"):
        styles[style].fontName = font
    styles["Title"].textColor = colors.HexColor("#B35C20")
    styles["Title"].fontSize = 18
    styles["Normal"].fontSize = 9
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 11
    def P(value, style="Normal"):
        return Paragraph(html.escape(str(value)).replace("\n", "<br/>"), styles[style])
    seller = quote["issuer_snapshot"]["seller"]
    client = quote["client"]
    story = [
        P("COTIZACIÓN COMERCIAL" if quote["status"] == "issued" else "BORRADOR DE COTIZACIÓN", "Title"),
        P("Folio: " + (quote["folio"] or "POR ASIGNAR (BORRADOR)")),
        P("Fecha: " + quote["quote_date"] + " | Válida hasta: " + quote["valid_until"]),
        Spacer(1, 8), P("EMISOR", "Heading2"),
        P(seller["company_name"] + " | " + seller["legal_name"]),
        P(seller["contact_name"] + " | " + seller["email"] + " | " + seller["phone"]),
        P(seller["address"]),
        Spacer(1, 8), P("CLIENTE", "Heading2"),
        P(client["name"] + (" | " + client["company_name"] if client["company_name"] else "")),
        P(client["email"] or ""),
        Spacer(1, 12), P("CONCEPTOS", "Heading2"),
    ]
    rows = [[P(x, "BodyText") for x in ("Descripción", "Cant.", "P. unit.", "Desc.", "Impuestos", "Total")]]
    for line in quote["items"]:
        rows.append([
            P(line["description"], "BodyText"), P(line["quantity"], "BodyText"),
            P(_money(line["unit_price"], quote), "BodyText"),
            P(_money(line["discount"], quote), "BodyText"),
            P("+" + _money(line["tax_added"], quote) + " / −" +
              _money(line["tax_withheld"], quote), "BodyText"),
            P(_money(line["line_total"], quote), "BodyText")
        ])
    table = Table(rows, repeatRows=1, hAlign="LEFT", colWidths=[155, 36, 65, 62, 86, 75])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F9EBDC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LINEBELOW", (0, 0), (-1, 0), .6, colors.HexColor("#B35C20")),
        ("LINEBELOW", (0, 1), (-1, -1), .15, colors.lightgrey),
    ]))
    story += [table, Spacer(1, 12), P("RESUMEN", "Heading2")]
    t = quote["totals"]
    for title, key in (
        ("Importe bruto", "gross"), ("Descuentos", "discount"),
        ("Base", "taxable_base"), ("Impuestos adicionales", "tax_added"),
        ("Retenciones", "tax_withheld"), ("TOTAL", "total"),
    ):
        story.append(P(title + ": " + _money(t[key], quote),
                       "Heading2" if key == "total" else "Normal"))
    for tax in quote["tax_breakdown"]:
        story.append(P(("+" if tax["kind"] == "add" else "−") +
                       tax["label"] + " (" + tax["code"] + "): " +
                       _money(tax["amount"], quote), "BodyText"))
    story += [
        Spacer(1, 10), P("CONDICIONES", "Heading2"),
        P("Pago: " + (quote["terms"]["payment_terms"] or "Por definir")),
        P("Entrega: " + (quote["terms"]["delivery_terms"] or "Por definir")),
        P("Notas: " + (quote["terms"]["notes"] or "Ninguna")),
    ]
    if quote["exclusions"]:
        story.append(P("EXCLUSIONES", "Heading2"))
        for exclusion in quote["exclusions"]:
            story.append(P(exclusion, "BodyText"))
    story.extend([
        Spacer(1, 10), P(quote["notice"], "BodyText"),
        P("No se ha enviado al cliente. Verifique los impuestos y condiciones.", "BodyText"),
    ])
    if quote["status"] == "issued":
        story.append(P("Aprobación declarada por: " + quote["approved_by"]))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=42, rightMargin=42,
                      topMargin=42, bottomMargin=42, title="MANGO Cotización").build(story)


def export_quote(quote, out_dir, formats=("json", "md"), allow_matching=False):
    preflight_formats(formats)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = quote["folio"] if quote["folio"] else "DRAFT-" + quote["draft_id"]
    paths = {fmt: out_dir / (stem + "." + fmt) for fmt in formats}
    exists = [fmt for fmt, path in paths.items() if path.exists()]
    if exists:
        # If partial output exists, do not mix old/new files. When re-running
        # an issued quote, an identical JSON sibling is the minimum guard.
        guard = out_dir / (stem + ".json")
        if not allow_matching or not guard.is_file():
            raise ValueError("Hay archivos de salida existentes; elige otro --out-dir.")
        actual = json.loads(guard.read_text(encoding="utf-8"))
        if actual.get("integrity") != quote.get("integrity") or actual != quote:
            raise ValueError("El JSON existente no coincide con la cotización aprobada.")
        if any(path.is_symlink() for path in paths.values()):
            raise ValueError("Enlaces simbólicos en salida: no permitidos.")
        if all(path.is_file() for path in paths.values()):
            return {k: str(v) for k, v in paths.items()}
        for path in paths.values():
            if path.exists() and path.is_symlink():
                raise ValueError("Archivo de salida simbólico no permitido.")
    for fmt, path in paths.items():
        if path.exists():
            continue
        # Reserve path exclusively before giving it to the renderer.
        fd = os.open(str(path), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        os.close(fd)
        try:
            if fmt == "json":
                path.write_text(json.dumps(quote, ensure_ascii=False, indent=2) + "\n",
                                encoding="utf-8")
            elif fmt == "md":
                path.write_text(markdown_quote(quote), encoding="utf-8")
            elif fmt == "docx":
                docx_quote(quote, path)
            else:
                pdf_quote(quote, path)
            try:
                path.chmod(0o600)
            except OSError:
                pass
        except BaseException:
            path.unlink(missing_ok=True)
            raise
    return {fmt: str(path) for fmt, path in paths.items()}
